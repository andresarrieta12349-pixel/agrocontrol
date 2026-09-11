"""
routers/reports.py
--------------------
Reportes descargables de AgroControl Pro:
  - Inventario actual (stock, valor, estado crítico)
  - Kardex de movimientos (entradas/salidas/ajustes por rango de fechas)
  - Cosechas (por ciclo productivo, opcionalmente filtradas)

Cada reporte está disponible en dos formatos: Excel (.xlsx) y PDF.
"""
import io
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import and_

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils.cell import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from app import models
from app.database import get_db
from app.auth import get_current_user

router = APIRouter(prefix="/api/reportes", tags=["Reportes"])

VERDE_BOSQUE = "1A4D2E"
VERDE_CLARO = "E4EFE0"


def _construir_excel(titulo: str, encabezados: list[str], filas: list[list]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(encabezados))
    celda_titulo = ws.cell(row=1, column=1, value=titulo)
    celda_titulo.font = Font(size=14, bold=True, color=VERDE_BOSQUE)
    ws.row_dimensions[1].height = 26

    fecha_celda = ws.cell(row=2, column=1, value=f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    fecha_celda.font = Font(size=9, italic=True, color="666666")

    fila_encabezado = 4
    for col_idx, encabezado in enumerate(encabezados, start=1):
        celda = ws.cell(row=fila_encabezado, column=col_idx, value=encabezado)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill(start_color=VERDE_BOSQUE, end_color=VERDE_BOSQUE, fill_type="solid")
        celda.alignment = Alignment(horizontal="center")

    for fila_idx, fila in enumerate(filas, start=fila_encabezado + 1):
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            if fila_idx % 2 == 0:
                celda.fill = PatternFill(start_color=VERDE_CLARO, end_color=VERDE_CLARO, fill_type="solid")

    for col_idx, encabezado in enumerate(encabezados, start=1):
        ancho = max(len(str(encabezado)), *(len(str(f[col_idx - 1])) for f in filas)) if filas else len(str(encabezado))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(ancho + 3, 12), 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _construir_pdf(titulo: str, encabezados: list[str], filas: list[list], subtitulo: str = "") -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    estilos = getSampleStyleSheet()

    estilo_titulo = ParagraphStyle(
        "TituloReporte", parent=estilos["Heading1"], textColor=colors.HexColor("#" + VERDE_BOSQUE), fontSize=18
    )
    estilo_meta = ParagraphStyle("Meta", parent=estilos["Normal"], textColor=colors.grey, fontSize=9)

    elementos = [
        Paragraph("AgroControl Pro", estilo_meta),
        Paragraph(titulo, estilo_titulo),
    ]
    if subtitulo:
        elementos.append(Paragraph(subtitulo, estilo_meta))
    elementos.append(Paragraph(f"Generado: {datetime.now().strftime('%Y-%m-%d %H:%M')}", estilo_meta))
    elementos.append(Spacer(1, 14))

    datos_tabla = [encabezados] + [[str(v) for v in fila] for fila in filas]
    if not filas:
        datos_tabla.append(["Sin registros para los filtros seleccionados"] + [""] * (len(encabezados) - 1))

    tabla = Table(datos_tabla, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + VERDE_BOSQUE)),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, -1), 8.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#" + VERDE_CLARO)]),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#dcdfd2")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    elementos.append(tabla)

    doc.build(elementos)
    buffer.seek(0)
    return buffer


def _respuesta_excel(buffer: io.BytesIO, nombre_archivo: str) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.xlsx"'},
    )


def _respuesta_pdf(buffer: io.BytesIO, nombre_archivo: str) -> StreamingResponse:
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}.pdf"'},
    )


def _obtener_filas_inventario(db: Session):
    productos = (
        db.query(models.Producto)
        .filter(models.Producto.activo.is_(True))
        .order_by(models.Producto.nombre)
        .all()
    )
    encabezados = ["Código", "Nombre", "Categoría", "Stock actual", "Unidad", "Stock mínimo", "Valor (COP)", "Estado"]
    filas = []
    for p in productos:
        valor = round((p.stock_actual or 0) * (p.precio_unitario or 0), 2)
        estado = "CRÍTICO" if p.en_stock_critico else "Normal"
        filas.append([p.codigo, p.nombre, p.categoria or "—", p.stock_actual, p.unidad_medida, p.stock_minimo, valor, estado])
    return encabezados, filas


@router.get("/inventario/excel", summary="Descargar reporte de inventario actual (Excel)")
def reporte_inventario_excel(db: Session = Depends(get_db), usuario_actual=Depends(get_current_user)):
    encabezados, filas = _obtener_filas_inventario(db)
    buffer = _construir_excel("Inventario actual", encabezados, filas)
    return _respuesta_excel(buffer, f"inventario_{date.today().isoformat()}")


@router.get("/inventario/pdf", summary="Descargar reporte de inventario actual (PDF)")
def reporte_inventario_pdf(db: Session = Depends(get_db), usuario_actual=Depends(get_current_user)):
    encabezados, filas = _obtener_filas_inventario(db)
    buffer = _construir_pdf("Reporte de Inventario Actual", encabezados, filas)
    return _respuesta_pdf(buffer, f"inventario_{date.today().isoformat()}")


def _obtener_filas_kardex(db: Session, desde: Optional[date], hasta: Optional[date]):
    query = db.query(models.MovimientoInventario)
    condiciones = []
    if desde:
        condiciones.append(models.MovimientoInventario.fecha >= datetime.combine(desde, datetime.min.time()))
    if hasta:
        condiciones.append(models.MovimientoInventario.fecha <= datetime.combine(hasta, datetime.max.time()))
    if condiciones:
        query = query.filter(and_(*condiciones))

    movimientos = query.order_by(models.MovimientoInventario.fecha.desc()).limit(2000).all()
    encabezados = ["Fecha", "Producto", "Bodega", "Tipo", "Cantidad", "Costo unitario", "Referencia"]
    filas = []
    for m in movimientos:
        filas.append([
            m.fecha.strftime("%Y-%m-%d %H:%M"),
            m.producto.nombre if m.producto else f"#{m.producto_id}",
            m.bodega.nombre if m.bodega else f"#{m.bodega_id}",
            m.tipo.value if hasattr(m.tipo, "value") else m.tipo,
            m.cantidad,
            m.costo_unitario or 0,
            m.referencia or "—",
        ])
    return encabezados, filas


@router.get("/kardex/excel", summary="Descargar kardex de movimientos (Excel)")
def reporte_kardex_excel(
    desde: Optional[date] = Query(None), hasta: Optional[date] = Query(None),
    db: Session = Depends(get_db), usuario_actual=Depends(get_current_user),
):
    encabezados, filas = _obtener_filas_kardex(db, desde, hasta)
    buffer = _construir_excel("Kardex de movimientos", encabezados, filas)
    return _respuesta_excel(buffer, f"kardex_{date.today().isoformat()}")


@router.get("/kardex/pdf", summary="Descargar kardex de movimientos (PDF)")
def reporte_kardex_pdf(
    desde: Optional[date] = Query(None), hasta: Optional[date] = Query(None),
    db: Session = Depends(get_db), usuario_actual=Depends(get_current_user),
):
    encabezados, filas = _obtener_filas_kardex(db, desde, hasta)
    subtitulo = f"Periodo: {desde or 'inicio'} — {hasta or 'hoy'}"
    buffer = _construir_pdf("Kardex de Movimientos de Inventario", encabezados, filas, subtitulo)
    return _respuesta_pdf(buffer, f"kardex_{date.today().isoformat()}")


def _obtener_filas_cosechas(db: Session, ciclo_id: Optional[int]):
    query = db.query(models.Cosecha)
    if ciclo_id:
        query = query.filter(models.Cosecha.ciclo_id == ciclo_id)
    cosechas = query.order_by(models.Cosecha.fecha.desc()).all()

    encabezados = ["Fecha", "Ciclo productivo", "Cultivo", "Cantidad (kg)", "Calidad", "Observaciones"]
    filas = []
    for c in cosechas:
        ciclo = c.ciclo
        filas.append([
            c.fecha.strftime("%Y-%m-%d"),
            ciclo.nombre if ciclo else f"#{c.ciclo_id}",
            ciclo.cultivo if ciclo else "—",
            c.cantidad_kg,
            c.calidad or "—",
            c.observaciones or "—",
        ])
    total_kg = sum(c.cantidad_kg for c in cosechas)
    return encabezados, filas, total_kg


@router.get("/cosechas/excel", summary="Descargar reporte de cosechas (Excel)")
def reporte_cosechas_excel(
    ciclo_id: Optional[int] = Query(None),
    db: Session = Depends(get_db), usuario_actual=Depends(get_current_user),
):
    encabezados, filas, total_kg = _obtener_filas_cosechas(db, ciclo_id)
    filas_con_total = filas + [["", "", "TOTAL", total_kg, "", ""]] if filas else filas
    buffer = _construir_excel("Cosechas", encabezados, filas_con_total)
    return _respuesta_excel(buffer, f"cosechas_{date.today().isoformat()}")


@router.get("/cosechas/pdf", summary="Descargar reporte de cosechas (PDF)")
def reporte_cosechas_pdf(
    ciclo_id: Optional[int] = Query(None),
    db: Session = Depends(get_db), usuario_actual=Depends(get_current_user),
):
    encabezados, filas, total_kg = _obtener_filas_cosechas(db, ciclo_id)
    subtitulo = f"Total cosechado: {total_kg} kg" + (f" — Ciclo #{ciclo_id}" if ciclo_id else " — Todos los ciclos")
    buffer = _construir_pdf("Reporte de Cosechas", encabezados, filas, subtitulo)
    return _respuesta_pdf(buffer, f"cosechas_{date.today().isoformat()}")